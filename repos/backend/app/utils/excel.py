"""
Excel 工具模块 (Excel Utilities)

此模块提供 Excel 文件的读写功能，用于学生数据的导入和模板生成

依赖库：
- openpyxl: Python 操作 Excel 文件的库（支持 .xlsx 格式）

功能说明：
1. read_student_import(): 读取学生导入的 Excel 文件
2. create_student_template(): 生成学生导入模板
"""

from io import BytesIO  # 内存文件对象，用于存储 Excel 字节数据
from openpyxl import load_workbook, Workbook  # Excel 文件读写


# ========== 读取学生导入文件 ==========

def read_student_import(file_content: bytes) -> list[tuple[str, str]]:
    """
    读取学生导入 Excel 文件

    从 Excel 文件中解析学生数据，返回姓名和学号列表

    执行流程：
    1. 将字节数据加载为 Workbook 对象
    2. 获取活动工作表
    3. 从第二行开始遍历（跳过表头）
    4. 读取每行的姓名和学号
    5. 返回 (姓名, 学号) 元组列表

    Excel 文件格式要求：
    - 第一行：表头（学生姓名、学生学号）- 会被跳过
    - 第二行起：数据行
    - 第一列：学生姓名
    - 第二列：学生学号

    参数：
        file_content: Excel 文件的字节内容（HTTP 上传的文件内容）

    返回：
        List[Tuple[str, str]]: 学生数据列表
        - 每个元素是 (姓名, 学号) 元组
        - 例如：[("张三", "2025001"), ("李四", "2025002")]

    异常：
        ValueError: 文件格式错误或无法解析
    """
    try:
        # Step 1: 将字节数据包装为 BytesIO 对象
        # BytesIO 允许我们像操作文件一样操作内存中的字节数据
        bytes_io = BytesIO(file_content)

        # Step 2: 使用 openpyxl 加载 Excel 工作簿
        workbook = load_workbook(bytes_io)

        # Step 3: 获取活动工作表（Excel 默认的 Sheet）
        sheet = workbook.active

        # Step 4: 准备存储学生数据的列表
        students: list[tuple[str, str]] = []

        # Step 5: 遍历工作表数据行（跳过表头，从第二行开始）
        # iter_rows() 迭代工作表的所有行
        # min_row=2: 从第二行开始（跳过表头）
        # max_col=2: 最多读取两列
        # values_only=True: 只返回值，不返回单元格对象
        for row in sheet.iter_rows(min_row=2, max_col=2, values_only=True):
            name, student_no = row

            # 跳过空行
            if name is not None and student_no is not None:
                # 去除首尾空白并转换为字符串
                students.append((str(name).strip(), str(student_no).strip()))

        return students

    except Exception as e:
        # 如果解析失败，抛出明确的错误信息
        raise ValueError(f"Failed to read Excel file: {e}")


# ========== 创建学生导入模板 ==========

def create_student_template() -> bytes:
    """
    创建学生导入模板 Excel 文件

    生成一个包含示例数据的模板文件，供管理员下载并填写学生信息

    执行流程：
    1. 创建新的 Workbook 和工作表
    2. 写入表头行（学生姓名、学生学号）
    3. 写入一行示例数据
    4. 调整列宽以适应内容
    5. 保存到 BytesIO 对象
    6. 返回字节数据

    返回：
        bytes: Excel 文件的字节内容，可直接用于 HTTP 响应下载

    Excel 文件内容：
    - 工作表名称：学生信息
    - 第一行（表头）：学生姓名 | 学生学号
    - 第二行（示例）：张三 | 001
    """
    # Step 1: 创建新的工作簿
    wb = Workbook()

    # Step 2: 获取活动工作表
    ws = wb.active

    # Step 3: 设置工作表名称
    ws.title = "学生信息"

    # Step 4: 写入表头行
    # enumerate(sequence, start=1) 同时返回索引和值
    # 设置 start=1 使列号从 1 开始（对应 Excel 的 A、B 列）
    headers = ["学生姓名", "学生学号"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)

    # Step 5: 写入示例数据行
    # 这一行作为参考，展示正确的格式
    ws.cell(row=2, column=1, value="张三")
    ws.cell(row=2, column=2, value="001")

    # Step 6: 调整列宽，使内容更易读
    ws.column_dimensions["A"].width = 15  # A 列（姓名）宽度 15
    ws.column_dimensions["B"].width = 15  # B 列（学号）宽度 15

    # Step 7: 保存到 BytesIO 对象
    output = BytesIO()
    wb.save(output)

    # Step 8: 将指针移到开头，准备被读取
    output.seek(0)

    # Step 9: 返回字节内容
    return output.getvalue()
